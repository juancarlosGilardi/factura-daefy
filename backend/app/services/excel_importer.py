"""Importadores Excel para clientes y productos.

Tolera variaciones de nombres de columnas (lowercase, sin acentos, sinónimos
comunes como "razon_social" / "nombres" / "nombre").

Devuelve siempre la tupla (creados, actualizados, errores: list[dict]) para
que el endpoint pueda renderizar la lista de errores fila por fila sin
abortar todo el archivo.
"""
from __future__ import annotations

import io
import re
import unicodedata
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session

from ..models.cliente import Cliente
from ..models.producto import Producto


# ---------------------------------------------------------------------------
# Utilidades comunes
# ---------------------------------------------------------------------------
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")


def _normaliza_columna(name: Any) -> str:
    """Normaliza un nombre de columna: minúsculas, sin tildes, sin espacios.

    Ej: "Razón Social" -> "razon_social"; " N° de Documento " -> "n_de_documento".
    """
    if name is None:
        return ""
    s = str(name).strip().lower()
    # quitar tildes
    s = "".join(
        c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)
    )
    # reemplazar separadores por _
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def _resolver_columna(cols_normalizadas: dict[str, str], *aliases: str) -> str | None:
    """Dado un dict {col_normalizada: col_original}, devuelve la col_original
    cuyo nombre normalizado coincida con cualquiera de los aliases (también
    se normalizan los aliases por si vienen con tildes/espacios)."""
    for alias in aliases:
        key = _normaliza_columna(alias)
        if key in cols_normalizadas:
            return cols_normalizadas[key]
    return None


def _cell_str(row: pd.Series, col: str | None) -> str:
    if col is None:
        return ""
    val = row.get(col, None)
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    # Bug E4: documentos numéricos venían como float (Excel los autotipa).
    # Para enteros formateamos sin decimales; para no-enteros usamos
    # representación "limpia" (sin notación científica) que la validación
    # posterior rechazará por contener punto.
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.10g}"
    return str(val).strip()


def _cell_float(row: pd.Series, col: str | None, default: float = 0.0) -> float:
    if col is None:
        return default
    val = row.get(col, None)
    if val is None or (isinstance(val, float) and pd.isna(val)) or val == "":
        return default
    try:
        return float(str(val).replace(",", "."))
    except (ValueError, TypeError):
        return default


def _cell_bool(row: pd.Series, col: str | None, default: bool = False) -> bool:
    if col is None:
        return default
    val = row.get(col, None)
    if val is None or (isinstance(val, float) and pd.isna(val)) or val == "":
        return default
    s = str(val).strip().lower()
    return s in ("1", "si", "sí", "s", "true", "verdadero", "yes", "y", "x")


def _email_valido(email: str) -> bool:
    return bool(email) and bool(EMAIL_RE.match(email))


def _leer_excel(file_bytes: bytes) -> pd.DataFrame:
    """Lee un xlsx (bytes) en DataFrame. Solo la primera hoja.

    Bug E7: usar `data_only=True` para que las fórmulas devuelvan su valor
    cacheado y no la cadena de la fórmula.

    Bug A4/G3: tolerar filas de título mergeadas (1 sola celda no vacía)
    al inicio. Detecta la primera fila con >= 2 celdas no vacías como header.
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.values)
    if not rows:
        return pd.DataFrame()
    # Saltar filas de título: la fila header debe tener >= 2 celdas no vacías
    header_idx = 0
    for i, row in enumerate(rows):
        non_empty = sum(
            1 for v in row if v is not None and str(v).strip()
        )
        if non_empty >= 2:
            header_idx = i
            break
    headers = [str(h) if h is not None else "" for h in rows[header_idx]]
    return pd.DataFrame(rows[header_idx + 1:], columns=headers)


# ---------------------------------------------------------------------------
# Clientes
# ---------------------------------------------------------------------------
class ImportadorClientes:
    """Importa maestro de clientes desde xlsx.

    Columnas reconocidas (con tolerancia de variaciones):
      - tipo_doc / tipo_documento  (1=DNI, 6=RUC; default según longitud)
      - numero_doc / numero_documento / nro_doc / documento
      - razon_social / nombres / nombre / cliente
      - direccion
      - email / correo
      - telefono / celular
      - ubigeo
      - distrito / provincia / departamento

    La clave única es (tipo_documento, numero_documento). Si ya existe se
    actualiza, si no, se crea.
    """

    COLUMNAS_PLANTILLA = [
        "tipo_doc",
        "numero_doc",
        "razon_social",
        "direccion",
        "email",
        "telefono",
        "ubigeo",
        "distrito",
        "provincia",
        "departamento",
    ]

    EJEMPLO = [
        "6",
        "20123456789",
        "MI CLIENTE EJEMPLO SAC",
        "Av. Javier Prado 123",
        "ventas@ejemplo.pe",
        "987654321",
        "150101",
        "Lima",
        "Lima",
        "Lima",
    ]

    def __init__(self, db: Session):
        self.db = db

    # ------------------------------------------------------------------
    def importar(self, file_bytes: bytes) -> tuple[int, int, list[dict]]:
        creados = 0
        actualizados = 0
        errores: list[dict] = []

        try:
            df = _leer_excel(file_bytes)
        except Exception as e:  # noqa: BLE001
            return 0, 0, [{"fila": 0, "error": f"No se pudo leer el archivo: {e}"}]

        cols_norm = {_normaliza_columna(c): c for c in df.columns}

        col_tipo = _resolver_columna(cols_norm, "tipo_doc", "tipo_documento", "tipo")
        col_num = _resolver_columna(
            cols_norm,
            "numero_doc",
            "numero_documento",
            "nro_doc",
            "nro_documento",
            "documento",
            "ruc",
            "dni",
        )
        col_razon = _resolver_columna(
            cols_norm, "razon_social", "razon", "nombres", "nombre", "cliente"
        )
        col_dir = _resolver_columna(cols_norm, "direccion", "domicilio")
        col_email = _resolver_columna(cols_norm, "email", "correo", "e_mail")
        col_tel = _resolver_columna(cols_norm, "telefono", "celular", "movil")
        col_ubi = _resolver_columna(cols_norm, "ubigeo")
        col_dis = _resolver_columna(cols_norm, "distrito")
        col_prov = _resolver_columna(cols_norm, "provincia")
        col_dep = _resolver_columna(cols_norm, "departamento", "region")

        if col_num is None:
            return 0, 0, [
                {"fila": 0, "error": "Falta la columna de número de documento."}
            ]
        if col_razon is None:
            return 0, 0, [
                {"fila": 0, "error": "Falta la columna de razón social / nombres."}
            ]

        for idx, row in df.iterrows():
            fila = int(idx) + 2  # +1 por header, +1 por base 1

            numero = _cell_str(row, col_num)[:20]
            razon = _cell_str(row, col_razon)[:200]

            if not numero and not razon:
                # fila vacía, ignorar
                continue
            if not numero:
                errores.append({"fila": fila, "error": "Documento vacío."})
                continue
            if not razon:
                errores.append({"fila": fila, "error": "Razón social / nombre vacío."})
                continue

            tipo = _cell_str(row, col_tipo)
            if not tipo:
                # Inferir por longitud
                if len(numero) == 11:
                    tipo = "6"
                elif len(numero) == 8:
                    tipo = "1"
                else:
                    tipo = "0"  # otros

            # Validaciones
            # Bug E4: rechazar documentos con caracteres no-numéricos para
            # tipos que la SUNAT exige numéricos (DNI=1, RUC=6).
            if tipo in ("1", "6") and numero and not numero.isdigit():
                errores.append({
                    "fila": fila,
                    "error": (
                        f"Documento debe ser numérico para tipo {tipo}: {numero}"
                    ),
                })
                continue
            if tipo == "6" and not (numero.isdigit() and len(numero) == 11):
                errores.append(
                    {"fila": fila, "error": f"RUC inválido (debe tener 11 dígitos): {numero}"}
                )
                continue
            if tipo == "1" and not (numero.isdigit() and len(numero) == 8):
                errores.append(
                    {"fila": fila, "error": f"DNI inválido (debe tener 8 dígitos): {numero}"}
                )
                continue

            email = _cell_str(row, col_email)
            if email and not _email_valido(email):
                errores.append(
                    {
                        "fila": fila,
                        "error": f"Email inválido (se importó sin email): {email}",
                    }
                )
                email = ""

            data = dict(
                tipo_documento=tipo,
                numero_documento=numero,
                razon_social=razon[:200],
                direccion=_cell_str(row, col_dir)[:300] or None,
                email=email or None,
                telefono=_cell_str(row, col_tel)[:40] or None,
                ubigeo=_cell_str(row, col_ubi)[:6] or None,
                distrito=_cell_str(row, col_dis)[:80] or None,
                provincia=_cell_str(row, col_prov)[:80] or None,
                departamento=_cell_str(row, col_dep)[:80] or None,
            )

            try:
                existente = (
                    self.db.query(Cliente)
                    .filter_by(tipo_documento=tipo, numero_documento=numero)
                    .first()
                )
                if existente:
                    for k, v in data.items():
                        if v is not None:
                            setattr(existente, k, v)
                    actualizados += 1
                else:
                    self.db.add(Cliente(**data))
                    creados += 1
            except Exception as e:  # noqa: BLE001
                errores.append({"fila": fila, "error": f"Error de BD: {e}"})

        try:
            self.db.commit()
        except Exception as e:  # noqa: BLE001
            self.db.rollback()
            errores.append({"fila": 0, "error": f"Error al guardar: {e}"})
            return 0, 0, errores

        return creados, actualizados, errores

    # ------------------------------------------------------------------
    @classmethod
    def plantilla_xlsx(cls) -> bytes:
        return _construir_plantilla(
            titulo="Plantilla de clientes — Factura-mdb",
            columnas=cls.COLUMNAS_PLANTILLA,
            ejemplo=cls.EJEMPLO,
            ayuda=[
                "tipo_doc: 6=RUC, 1=DNI, 4=Carnet ext., 7=Pasaporte, 0=Otros",
                "numero_doc: solo dígitos. RUC=11, DNI=8.",
                "Si dejas tipo_doc vacío se infiere por la longitud del documento.",
            ],
        )


# ---------------------------------------------------------------------------
# Productos
# ---------------------------------------------------------------------------
class ImportadorProductos:
    """Importa maestro de productos desde xlsx.

    Columnas reconocidas:
      - codigo / cod / sku
      - descripcion / nombre / producto
      - unidad_medida / unidad / um  (default NIU)
      - precio_unitario / valor_unitario / precio
      - moneda  (PEN/USD, default PEN)
      - afectacion_igv / tipo_afectacion_igv  (cat 07, default 10=gravado)
      - incluye_igv  (true/false, default false)
      - notas
    """

    COLUMNAS_PLANTILLA = [
        "codigo",
        "descripcion",
        "unidad_medida",
        "precio_unitario",
        "moneda",
        "afectacion_igv",
        "incluye_igv",
        "notas",
    ]

    EJEMPLO = [
        "PROD001",
        "Producto de ejemplo",
        "NIU",
        "100.00",
        "PEN",
        "10",
        "false",
        "",
    ]

    AFECTACIONES_VALIDAS = {
        "10", "11", "12", "13", "14", "15", "16", "17",  # gravadas
        "20", "21",  # exoneradas
        "30", "31", "32", "33", "34", "35", "36", "37",  # inafectas
        "40",  # exportación
    }

    def __init__(self, db: Session):
        self.db = db

    # ------------------------------------------------------------------
    def importar(self, file_bytes: bytes) -> tuple[int, int, list[dict]]:
        creados = 0
        actualizados = 0
        errores: list[dict] = []

        try:
            df = _leer_excel(file_bytes)
        except Exception as e:  # noqa: BLE001
            return 0, 0, [{"fila": 0, "error": f"No se pudo leer el archivo: {e}"}]

        cols_norm = {_normaliza_columna(c): c for c in df.columns}

        col_cod = _resolver_columna(cols_norm, "codigo", "cod", "sku")
        col_desc = _resolver_columna(cols_norm, "descripcion", "nombre", "producto")
        col_um = _resolver_columna(cols_norm, "unidad_medida", "unidad", "um")
        col_precio = _resolver_columna(
            cols_norm,
            "precio_unitario",
            "valor_unitario",
            "precio",
            "valor",
        )
        col_moneda = _resolver_columna(cols_norm, "moneda", "currency")
        col_afect = _resolver_columna(
            cols_norm, "afectacion_igv", "tipo_afectacion_igv", "afectacion", "igv"
        )
        col_incluye = _resolver_columna(
            cols_norm, "incluye_igv", "con_igv", "precio_con_igv"
        )
        col_notas = _resolver_columna(cols_norm, "notas", "observaciones", "nota")

        if col_cod is None:
            return 0, 0, [{"fila": 0, "error": "Falta la columna 'codigo'."}]
        if col_desc is None:
            return 0, 0, [{"fila": 0, "error": "Falta la columna 'descripcion'."}]

        for idx, row in df.iterrows():
            fila = int(idx) + 2

            codigo = _cell_str(row, col_cod)[:40]
            descripcion = _cell_str(row, col_desc)

            if not codigo and not descripcion:
                continue
            if not codigo:
                errores.append({"fila": fila, "error": "Código vacío."})
                continue
            if not descripcion:
                errores.append({"fila": fila, "error": "Descripción vacía."})
                continue

            unidad = _cell_str(row, col_um) or "NIU"
            precio = _cell_float(row, col_precio, 0.0)
            moneda = (_cell_str(row, col_moneda) or "PEN").upper()[:3]
            afect = _cell_str(row, col_afect) or "10"
            if afect not in self.AFECTACIONES_VALIDAS:
                errores.append(
                    {
                        "fila": fila,
                        "error": (
                            f"Afectación IGV '{afect}' inválida, se usó 10 (gravado)."
                        ),
                    }
                )
                afect = "10"
            incluye = _cell_bool(row, col_incluye, False)

            data = dict(
                codigo=codigo[:40],
                descripcion=descripcion[:300],
                unidad_medida=unidad[:8],
                valor_unitario=precio,
                moneda=moneda,
                tipo_afectacion_igv=afect,
                incluye_igv=incluye,
                notas=_cell_str(row, col_notas)[:500] or None,
            )

            try:
                existente = self.db.query(Producto).filter_by(codigo=codigo).first()
                if existente:
                    for k, v in data.items():
                        setattr(existente, k, v)
                    actualizados += 1
                else:
                    self.db.add(Producto(**data))
                    creados += 1
            except Exception as e:  # noqa: BLE001
                errores.append({"fila": fila, "error": f"Error de BD: {e}"})

        try:
            self.db.commit()
        except Exception as e:  # noqa: BLE001
            self.db.rollback()
            errores.append({"fila": 0, "error": f"Error al guardar: {e}"})
            return 0, 0, errores

        return creados, actualizados, errores

    @classmethod
    def plantilla_xlsx(cls) -> bytes:
        return _construir_plantilla(
            titulo="Plantilla de productos — Factura-mdb",
            columnas=cls.COLUMNAS_PLANTILLA,
            ejemplo=cls.EJEMPLO,
            ayuda=[
                "unidad_medida: catálogo SUNAT 03 (NIU=unidad, ZZ=servicio, KGM=kg, ...)",
                "afectacion_igv: 10=Gravado, 20=Exonerado, 30=Inafecto, 40=Exportación",
                "incluye_igv: true si el precio ya tiene IGV; false si es valor sin IGV.",
            ],
        )


# ---------------------------------------------------------------------------
# Plantilla helper
# ---------------------------------------------------------------------------
def _construir_plantilla(
    titulo: str,
    columnas: Iterable[str],
    ejemplo: Iterable[str],
    ayuda: Iterable[str],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    azul = "FF004D98"
    grana = "FFA50044"
    dorado_bg = "FFFFF6CC"

    # Fila 1: título
    ws.cell(row=1, column=1, value=titulo)
    ws.cell(row=1, column=1).font = Font(bold=True, color="FFFFFFFF", size=12)
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor=azul)
    cols = list(columnas)
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=max(1, len(cols))
    )
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Fila 2: headers
    for i, name in enumerate(cols, start=1):
        c = ws.cell(row=2, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=grana)
        c.alignment = Alignment(horizontal="center")

    # Fila 3: ejemplo
    for i, val in enumerate(ejemplo, start=1):
        c = ws.cell(row=3, column=i, value=val)
        c.fill = PatternFill("solid", fgColor=dorado_bg)

    # Anchos cómodos
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = max(
            14, len(str(name)) + 4
        )

    # Hoja "Ayuda"
    ws2 = wb.create_sheet("Ayuda")
    ws2.cell(row=1, column=1, value="Notas de uso").font = Font(bold=True, size=12)
    for i, linea in enumerate(ayuda, start=2):
        ws2.cell(row=i, column=1, value=f"• {linea}")
    ws2.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
