"""Generador de reportes Excel formateados para impresión A4.

Este módulo se enfoca en producir XLSX con configuración de página
optimizada para impresión vertical en A4: márgenes pequeños, anchos
ajustados al contenido, headers repetidos en cada página, y un bloque
de resumen al final.

Diseñado para el reporte de ventas del día — pero los helpers de estilo
son reutilizables para otros reportes imprimibles.
"""
from __future__ import annotations

import io
from datetime import date as _date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------

_THIN = Side(style="thin", color="FF999999")
_BORDER_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_HEADER_FILL = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9",
                            fill_type="solid")
_ALT_FILL = PatternFill(start_color="FFF7F7F7", end_color="FFF7F7F7",
                         fill_type="solid")

_TITLE_FONT = Font(name="Calibri", bold=True, size=14)
_SUBTITLE_FONT = Font(name="Calibri", bold=True, size=11)
_HEADER_FONT = Font(name="Calibri", bold=True, size=10)
_DATA_FONT = Font(name="Calibri", size=10)
_RESUMEN_FONT = Font(name="Calibri", bold=True, size=12)

_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _format_fecha_es(d: _date) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year:04d}"


def _format_pen(v: float | int) -> str:
    try:
        return f"S/ {float(v):,.2f}"
    except (TypeError, ValueError):
        return "S/ 0.00"


def _setup_a4_print(ws: Worksheet, header_rows: int = 3) -> None:
    """Configura la hoja para impresión A4 vertical con márgenes ajustados.

    Repite las primeras `header_rows` filas como título de impresión en
    cada página y ajusta al ancho de la hoja (alto auto).
    """
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # múltiples páginas verticalmente
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins = PageMargins(
        left=0.4, right=0.4, top=0.5, bottom=0.5,
        header=0.2, footer=0.2,
    )

    # Repetir las primeras `header_rows` filas en cada página impresa
    ws.print_title_rows = f"1:{header_rows}"

    # Footer con número de página y fecha de impresión
    ws.oddFooter.left.text = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws.oddFooter.left.size = 9
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.right.size = 9


def _autosize_columns(
    ws: Worksheet,
    headers: Sequence[tuple[str, str]],
    rows: Sequence[dict],
    min_w: int = 8,
    max_w: int = 40,
) -> None:
    """Ajusta el ancho de cada columna al contenido (con tope)."""
    for col_idx, (key, label) in enumerate(headers, start=1):
        max_len = len(label)
        for r in rows:
            v = r.get(key)
            if v is None:
                continue
            length = len(str(v))
            if length > max_len:
                max_len = length
        width = min(max(max_len + 2, min_w), max_w)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


# ---------------------------------------------------------------------------
# Reporte de ventas del día
# ---------------------------------------------------------------------------

# Columnas (key, label, alineación: 'l'/'r'/'c', es_numero)
_COLS_VENTAS_DIA: list[tuple[str, str, str, bool]] = [
    ("tipo_doc", "Tipo", "l", False),
    ("serie_numero", "Serie-Número", "l", False),
    ("hora", "Hora", "c", False),
    ("cliente", "Cliente", "l", False),
    ("ruc_dni", "RUC/DNI", "l", False),
    ("moneda", "Mon", "c", False),
    ("subtotal", "Subtotal", "r", True),
    ("igv", "IGV", "r", True),
    ("total", "Total", "r", True),
    ("estado", "Estado", "c", False),
]


def build_ventas_dia_xlsx(
    fecha: _date,
    items: Sequence[dict],
    resumen: dict,
    empresa_nombre: str = "DAEFY S.A.C.",
    empresa_ruc: str = "20615413071",
    logo_path: str | Path | None = None,
) -> bytes:
    """Genera el XLSX del reporte de ventas del día listo para impresión A4.

    Estructura:
      Filas 1-3: encabezado (logo + empresa, título con fecha, separador).
                 Se repite en cada página al imprimir.
      Fila 4:    headers de tabla (gris claro, negrita, bordes).
      Filas 5-N: datos (alternadas, bordes finos, números a la derecha).
      Fila N+1:  línea separadora.
      Fila N+2:  bloque de resumen (negrita, font 12).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas del día"

    n_cols = len(_COLS_VENTAS_DIA)
    last_col = get_column_letter(n_cols)

    # ------- Fila 1: empresa (con logo si existe) -------
    cell_empresa = ws.cell(
        row=1, column=1, value=f"{empresa_nombre} - RUC {empresa_ruc}",
    )
    cell_empresa.font = Font(name="Calibri", bold=True, size=11)
    cell_empresa.alignment = _CENTER
    ws.merge_cells(f"A1:{last_col}1")
    ws.row_dimensions[1].height = 24

    if logo_path:
        try:
            p = Path(logo_path)
            if p.exists():
                img = XLImage(str(p))
                # Limitar tamaño del logo a ~60x60 px
                img.width = 60
                img.height = 60
                ws.add_image(img, "A1")
                ws.row_dimensions[1].height = 50
        except Exception:  # noqa: BLE001
            # El logo es decorativo — si openpyxl no lo soporta, seguimos.
            pass

    # ------- Fila 2: título con fecha -------
    cell_titulo = ws.cell(
        row=2, column=1,
        value=f"REPORTE DE VENTAS — {_format_fecha_es(fecha)}",
    )
    cell_titulo.font = _TITLE_FONT
    cell_titulo.alignment = _CENTER
    ws.merge_cells(f"A2:{last_col}2")
    ws.row_dimensions[2].height = 22

    # ------- Fila 3: línea separadora (borde inferior) -------
    sep_border = Border(bottom=Side(style="medium", color="FF000000"))
    for c in range(1, n_cols + 1):
        ws.cell(row=3, column=c).border = sep_border
    ws.row_dimensions[3].height = 6

    # ------- Fila 4: headers de tabla -------
    header_row = 4
    for col_idx, (_, label, _align, _is_num) in enumerate(_COLS_VENTAS_DIA, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER_ALL
    ws.row_dimensions[header_row].height = 20

    # ------- Filas 5..N: datos -------
    data_start = header_row + 1
    row_idx = data_start
    for i, item in enumerate(items):
        for col_idx, (key, _label, align, is_num) in enumerate(_COLS_VENTAS_DIA, start=1):
            val = item.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _DATA_FONT
            if align == "r":
                cell.alignment = _RIGHT
            elif align == "c":
                cell.alignment = _CENTER
            else:
                cell.alignment = _LEFT
            cell.border = _BORDER_ALL
            if is_num and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
            # Filas alternadas
            if i % 2 == 1:
                cell.fill = _ALT_FILL
        row_idx += 1

    # Si no hay items, fila vacía con mensaje
    if not items:
        cell = ws.cell(row=row_idx, column=1, value="(Sin ventas registradas para esta fecha)")
        cell.alignment = _CENTER
        cell.font = Font(name="Calibri", italic=True, color="FF777777")
        ws.merge_cells(start_row=row_idx, start_column=1,
                        end_row=row_idx, end_column=n_cols)
        row_idx += 1

    # ------- Línea separadora -------
    sep_row = row_idx
    for c in range(1, n_cols + 1):
        ws.cell(row=sep_row, column=c).border = sep_border
    ws.row_dimensions[sep_row].height = 6
    row_idx += 1

    # ------- Resumen -------
    resumen_text = (
        f"Cantidad Facturas: {resumen.get('cantidad_facturas', 0)}   |   "
        f"Cantidad Boletas: {resumen.get('cantidad_boletas', 0)}   |   "
        f"Total General: {_format_pen(resumen.get('total_general_pen', 0.0))}"
    )
    cell_res = ws.cell(row=row_idx, column=1, value=resumen_text)
    cell_res.font = _RESUMEN_FONT
    cell_res.alignment = _CENTER
    ws.merge_cells(start_row=row_idx, start_column=1,
                    end_row=row_idx, end_column=n_cols)
    ws.row_dimensions[row_idx].height = 24
    row_idx += 1

    # Segunda línea de resumen — desglose IGV
    detalle_text = (
        f"Total Facturas: {_format_pen(resumen.get('total_facturas_pen', 0.0))}   |   "
        f"Total Boletas: {_format_pen(resumen.get('total_boletas_pen', 0.0))}   |   "
        f"IGV total: {_format_pen(resumen.get('total_igv_pen', 0.0))}"
    )
    cell_det = ws.cell(row=row_idx, column=1, value=detalle_text)
    cell_det.font = _SUBTITLE_FONT
    cell_det.alignment = _CENTER
    ws.merge_cells(start_row=row_idx, start_column=1,
                    end_row=row_idx, end_column=n_cols)
    ws.row_dimensions[row_idx].height = 18

    # ------- Auto ancho columnas (basado en contenido + headers) -------
    headers_simple = [(k, l) for k, l, _, _ in _COLS_VENTAS_DIA]
    _autosize_columns(ws, headers_simple, items)

    # ------- Configuración de impresión A4 -------
    _setup_a4_print(ws, header_rows=4)  # filas 1-4 (logo+título+sep+headers)

    # ------- Serializar -------
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
