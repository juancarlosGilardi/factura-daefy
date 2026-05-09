"""Helpers para exportar datos a Excel (.xlsx) usando openpyxl.

Estilo blaugrana: header `#004D98` (azul) con texto blanco en negrita,
freeze_panes='A2', auto-ancho de columnas.

Uso típico::

    from .services.excel_export import dict_list_to_xlsx_bytes
    data = dict_list_to_xlsx_bytes(rows, headers=[("id", "ID"), ...], sheet_name="Ventas")
    return Response(content=data, media_type=XLSX_MEDIA_TYPE, headers={...})
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# Branding blaugrana
_HEADER_FILL = PatternFill(start_color="FF004D98", end_color="FF004D98",
                            fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _coerce(value: Any) -> Any:
    """Convierte tipos que openpyxl no maneja directamente."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, int, float, str, bool)):
        return value
    # Listas / dicts / etc. → str
    return str(value)


def _autosize(ws, num_cols: int, num_rows: int, max_width: int = 60) -> None:
    """Calcula ancho aproximado por columna basado en contenido."""
    for col_idx in range(1, num_cols + 1):
        max_len = 0
        letter = get_column_letter(col_idx)
        for row_idx in range(1, num_rows + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            length = len(str(val))
            if length > max_len:
                max_len = length
        # Padding y máximo razonable
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)


def _apply_header_style(ws, num_cols: int) -> None:
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def dict_list_to_xlsx_bytes(
    rows: Iterable[dict],
    headers: Sequence[tuple[str, str]] | Sequence[str],
    sheet_name: str = "Datos",
) -> bytes:
    """Convierte una lista de dicts a bytes .xlsx.

    Args:
        rows: iterable de dicts.
        headers: lista de (key, label) o lista de strings (key == label).
        sheet_name: nombre de la hoja (truncado a 31 chars por límite Excel).
    """
    # Normalizar headers a [(key, label), ...]
    norm_headers: list[tuple[str, str]] = []
    for h in headers:
        if isinstance(h, tuple):
            norm_headers.append((h[0], h[1]))
        else:
            norm_headers.append((str(h), str(h)))

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Datos")[:31]

    # Header row
    for col_idx, (_, label) in enumerate(norm_headers, start=1):
        ws.cell(row=1, column=col_idx, value=label)
    _apply_header_style(ws, len(norm_headers))

    # Data rows
    row_idx = 2
    for row in rows:
        for col_idx, (key, _) in enumerate(norm_headers, start=1):
            val = row.get(key) if isinstance(row, dict) else getattr(row, key, None)
            ws.cell(row=row_idx, column=col_idx, value=_coerce(val))
        row_idx += 1

    _autosize(ws, len(norm_headers), row_idx - 1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dataframe_to_xlsx_bytes(df, sheet_name: str = "Datos") -> bytes:
    """Convierte un pandas DataFrame a bytes .xlsx con estilo blaugrana.

    pandas es opcional: si no está instalado, este helper levanta ImportError.
    """
    try:
        import pandas as pd  # noqa: F401
    except Exception as e:  # noqa: BLE001
        raise ImportError("pandas requerido para dataframe_to_xlsx_bytes") from e

    headers = [(c, str(c)) for c in df.columns]
    rows = df.to_dict(orient="records")
    return dict_list_to_xlsx_bytes(rows, headers=headers, sheet_name=sheet_name)


def xlsx_response_headers(filename: str) -> dict[str, str]:
    """Headers HTTP estándar para descarga .xlsx."""
    safe = filename.replace('"', "")
    return {"Content-Disposition": f'attachment; filename="{safe}"'}
